"""Frozen synthetic hold-out. Never changes detectors in response to results."""

import argparse
from collections import Counter
from datetime import datetime
import hashlib
import io
import json
from pathlib import Path
import sys
import zipfile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))
from app.schemas.domain import Rule, Surface, ParsedDocument
from app.detectors.engine import inspect_rule
from app.parsers.document import parse_document
from app.services.coverage import rule_coverage
from app.services.ocr import provider
from app.services.vision import prepare_visual
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as SlideInches
from PIL import Image, ImageDraw, ImageFont
from lxml import etree
import numpy as np
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont

DEST = ROOT / "tests/generated/holdout"
NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


class NoVisionModel:
    configured = False


def font(size):
    for path in (
        "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ):
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def picture(kind):
    im = Image.new("RGB", (1100, 620), "#faf8f1")
    d = ImageDraw.Draw(im)
    texts = {
        "image_identity": "Author affiliation: Tidal Synthetic University",
        "ai_watermark": "Made with AI - synthetic illustration",
        "garbled": "Result: �?�?�?�?",
        "missing_glyph": "Labels: □□□□ □□□□",
        "replacement": "Terms: �����",
        "small_font": "Academic labels: variance 0.0082; alpha + beta; n = 2048",
    }
    if kind in texts:
        d.text(
            (35, 260),
            texts[kind],
            font=font(17 if kind == "small_font" else 39),
            fill="#514c44",
        )
    elif kind == "chart":
        d.line((100, 60, 100, 510, 960, 510), fill="#73684b", width=3)
        d.line(
            [(100 + i * 80, 470 - int(260 * (i / 10) ** 0.8)) for i in range(11)],
            fill="#897448",
            width=4,
        )
        d.text(
            (140, 535),
            "Time / measured response / synthetic chart",
            font=font(18),
            fill="#514c44",
        )
    elif kind == "badge":
        d.ellipse((380, 100, 630, 350), outline="#907545", width=13)
        d.ellipse((399, 119, 611, 331), outline="#907545", width=5)
        d.polygon(((433, 278), (507, 157), (580, 278)), fill="#907545")
        d.text((315, 420), "Fictional campus emblem", font=font(22), fill="#514c44")
    elif kind == "circle":
        d.ellipse((380, 100, 630, 350), outline="#cacbc8", width=3)
    elif kind == "qr_pattern":
        rng = np.random.default_rng(8291)
        for y in range(25):
            for x in range(25):
                if rng.integers(2):
                    d.rectangle(
                        (350 + x * 10, 140 + y * 10, 359 + x * 10, 149 + y * 10),
                        fill="#45443f",
                    )
    elif kind == "flow":
        for x, label in ((75, "Input"), (400, "Compute"), (725, "Output")):
            d.rounded_rectangle(
                (x, 235, x + 220, 340), radius=12, outline="#8d8170", width=3
            )
            d.text((x + 30, 270), label, font=font(25), fill="#514c44")
        d.line((295, 287, 400, 287), fill="#8d8170", width=3)
        d.line((620, 287, 725, 287), fill="#8d8170", width=3)
    elif kind == "apparatus":
        # Original apparatus illustration, not a real photograph.
        for x, y, w, h in ((140, 195, 240, 230), (640, 130, 280, 270)):
            d.rectangle(
                (x, y, x + w, y + h), fill="#deddd6", outline="#8c887e", width=3
            )
            d.rectangle((x + 30, y + 30, x + w - 30, y + 100), fill="#efeee9")
        d.line((380, 310, 520, 310, 520, 230, 640, 230), fill="#8c887e", width=8)
    else:
        raise ValueError("Unknown synthetic image kind")
    return im


def office_residue(path, case):
    with zipfile.ZipFile(path) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    xml = etree.fromstring(parts["word/document.xml"])
    p = xml.find(".//{%s}p" % NS)
    if case["kind"] == "revision":
        ins = etree.SubElement(
            p,
            "{%s}ins" % NS,
            {
                "{%s}id" % NS: "8",
                "{%s}author" % NS: "Fictional reviewer",
                "{%s}date" % NS: "2020-01-01T00:00:00Z",
            },
        )
        run = etree.SubElement(ins, "{%s}r" % NS)
        etree.SubElement(run, "{%s}t" % NS).text = case["text"]
    elif case["kind"] == "comment":
        comments = etree.Element("{%s}comments" % NS, nsmap={"w": NS})
        comment = etree.SubElement(
            comments,
            "{%s}comment" % NS,
            {"{%s}id" % NS: "2", "{%s}author" % NS: "Fictional reviewer"},
        )
        cp = etree.SubElement(comment, "{%s}p" % NS)
        run = etree.SubElement(cp, "{%s}r" % NS)
        etree.SubElement(run, "{%s}t" % NS).text = case["text"]
        parts["word/comments.xml"] = etree.tostring(
            comments, xml_declaration=True, encoding="UTF-8"
        )
        etree.SubElement(p, "{%s}commentRangeStart" % NS, {"{%s}id" % NS: "2"})
        etree.SubElement(p, "{%s}commentRangeEnd" % NS, {"{%s}id" % NS: "2"})
        rr = etree.SubElement(p, "{%s}r" % NS)
        etree.SubElement(rr, "{%s}commentReference" % NS, {"{%s}id" % NS: "2"})
        rel = etree.fromstring(parts["word/_rels/document.xml.rels"])
        etree.SubElement(
            rel,
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
            Id="rIdHoldoutComment",
            Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments",
            Target="comments.xml",
        )
        parts["word/_rels/document.xml.rels"] = etree.tostring(
            rel, xml_declaration=True, encoding="UTF-8"
        )
        types = etree.fromstring(parts["[Content_Types].xml"])
        etree.SubElement(
            types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
            PartName="/word/comments.xml",
            ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
        )
        parts["[Content_Types].xml"] = etree.tostring(
            types, xml_declaration=True, encoding="UTF-8"
        )
    parts["word/document.xml"] = etree.tostring(
        xml, xml_declaration=True, encoding="UTF-8"
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, content in parts.items():
            z.writestr(name, content)


def create_document(case):
    path = DEST / (case["id"] + "." + case["format"])
    kind = case["kind"]
    has_image = kind in {
        "image_identity",
        "ai_watermark",
        "garbled",
        "missing_glyph",
        "replacement",
        "small_font",
        "chart",
        "badge",
        "circle",
        "qr_pattern",
        "flow",
        "apparatus",
    }
    buf = io.BytesIO()
    if has_image:
        picture(kind).save(buf, format="PNG")
        buf.seek(0)
    if case["format"] in ("txt", "md"):
        path.write_text(case["text"], encoding="utf-8")
    elif case["format"] == "pdf":
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        c = canvas.Canvas(str(path), pagesize=(595, 842), invariant=1)
        c.setAuthor("FictionalHoldoutAuthor" if kind == "metadata" else "")
        c.setCreator(
            "" if kind == "normal_metadata" else "Synthetic document generator"
        )
        c.setFont("STSong-Light", 12)
        if has_image:
            c.drawImage(ImageReader(buf), 45, 230, width=495, height=279)
        else:
            for index, offset in enumerate(range(0, len(case["text"]), 38)):
                c.drawString(45, 690 - index * 20, case["text"][offset : offset + 38])
        c.save()
    elif case["format"] == "docx":
        doc = Document()
        doc.core_properties.author = (
            "FictionalHoldoutAuthor" if kind == "metadata" else ""
        )
        doc.core_properties.last_modified_by = ""
        doc.core_properties.created = doc.core_properties.modified = datetime(
            2020, 1, 1
        )
        p = doc.add_paragraph()
        if has_image:
            p.add_run().add_picture(buf, width=Inches(6))
        elif kind == "hidden":
            p.add_run(case["text"]).font.hidden = True
        elif kind not in ("comment", "revision"):
            p.add_run(case["text"])
        else:
            p.add_run("Synthetic anonymous material.")
        doc.save(path)
        if kind in ("comment", "revision"):
            office_residue(path, case)
    elif case["format"] == "pptx":
        deck = Presentation()
        deck.core_properties.author = (
            "FictionalHoldoutAuthor" if kind == "metadata" else ""
        )
        deck.core_properties.last_modified_by = ""
        deck.core_properties.created = deck.core_properties.modified = datetime(
            2020, 1, 1
        )
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        if has_image:
            slide.shapes.add_picture(
                buf, SlideInches(0.8), SlideInches(1), width=SlideInches(8)
            )
        else:
            box = slide.shapes.add_textbox(
                SlideInches(0.8), SlideInches(1), SlideInches(8), SlideInches(3)
            )
            box.text_frame.text = case["text"]
        if kind == "hidden_slide":
            slide._element.set("show", "0")
        deck.save(path)
    else:
        raise ValueError("Unknown document format")
    return path


def rule_for(case):
    return Rule(
        id="heldout-rule",
        source_rule_set_id="synthetic-heldout-confirmed",
        target=case["target"],
        category="原创合成评测",
        description="仅检查本条合成规则",
        scope=case["scope"],
        detection_method=case["method"],
    )


def inspect_case(case, parsed, path=None):
    if path and parsed.image_jobs:
        provider.inspect(path, parsed, lambda *args: None)
    rule = rule_for(case)
    if path and rule.detection_method == "visual":
        prepare_visual(path, parsed, model=NoVisionModel())
    findings = inspect_rule(rule, parsed, ["body", "metadata", "hidden", "images"])
    coverage = rule_coverage(
        rule, parsed, ["body", "metadata", "hidden", "images"], findings
    )
    evidence = [
        f for f in findings if f.get("surface_id") and f["status"] in ("FAIL", "REVIEW")
    ]
    verified_pass = (
        bool(findings)
        and all(f["status"] == "PASS" for f in findings)
        and coverage["status"] == "VERIFIED"
        and not parsed.warnings
    )
    return dict(
        id=case["id"],
        format=case.get("format", "sentence"),
        kind=case["kind"],
        target=case["target"],
        expected_risk=case["expected_risk"],
        candidate=bool(evidence),
        confirmed_fail=any(f["status"] == "FAIL" for f in evidence),
        actual_pass=verified_pass,
        rule_pass=any(f["status"] == "PASS" for f in findings),
        coverage=coverage["status"],
        statuses=[f["status"] for f in findings],
        diagnostics=parsed.warnings,
        evidence=[
            {
                k: f.get(k)
                for k in ("status", "evidence", "source_type", "page", "detector")
            }
            for f in evidence
        ],
        ocr_regions=sum(m["text_regions"] for m in parsed.ocr_metrics),
    )


def metrics(rows, field):
    tp = sum(r["expected_risk"] and r[field] for r in rows)
    fp = sum(not r["expected_risk"] and r[field] for r in rows)
    fn = sum(r["expected_risk"] and not r[field] for r in rows)
    tn = sum(not r["expected_risk"] and not r[field] for r in rows)
    p, r = tp / max(tp + fp, 1), tp / max(tp + fn, 1)
    positives = sum(row["expected_risk"] for row in rows)
    false_pass = sum(row["expected_risk"] and row["actual_pass"] for row in rows)
    return dict(
        tp=tp,
        fp=fp,
        fn=fn,
        tn=tn,
        precision=p,
        recall=r,
        f1=2 * p * r / max(p + r, 1e-12),
        false_positive=fp,
        false_negative=fn,
        false_pass=false_pass,
        false_pass_rate=false_pass / max(positives, 1),
        positive_rule_pass_count=sum(
            row["expected_risk"] and row["rule_pass"] for row in rows
        ),
    )


def summarize(rows):
    return dict(
        count=len(rows),
        evidence_candidate=metrics(rows, "candidate"),
        confirmed_fail=metrics(rows, "confirmed_fail"),
        coverage=dict(Counter(r["coverage"] for r in rows)),
    )


def validate_freeze(dataset_path, data):
    expected = dataset_path.with_suffix(".sha256").read_text().strip()
    if not frozen_bytes_match(dataset_path.read_bytes(), expected):
        raise ValueError("Dataset changed after freezing")
    for filename, digest in data["detector_file_hashes"].items():
        if not frozen_bytes_match((ROOT / filename).read_bytes(), digest):
            raise ValueError("Frozen detector changed: " + filename)
    return expected


def frozen_bytes_match(content, digest):
    # Git enforces LF; the original freeze used Windows uniform CRLF/LF files.
    # Accept only equivalent line endings, never regenerated semantic content.
    lf = content.replace(b"\r\n", b"\n")
    return digest in {
        hashlib.sha256(value).hexdigest()
        for value in (content, lf, lf.replace(b"\n", b"\r\n"))
    }


def main():
    cli = argparse.ArgumentParser()
    cli.add_argument("--smoke", action="store_true")
    cli.add_argument(
        "--output", type=Path, default=Path(__file__).with_name("results.json")
    )
    args = cli.parse_args()
    dataset_path = Path(__file__).with_name("dataset.json")
    data = json.loads(dataset_path.read_text(encoding="utf-8"))
    digest = validate_freeze(dataset_path, data)
    DEST.mkdir(parents=True, exist_ok=True)
    sentence_rows, document_rows = [], []
    sentences = data["sentence_cases"][:12] if args.smoke else data["sentence_cases"]
    documents = data["document_cases"][:5] if args.smoke else data["document_cases"]
    for case in sentences:
        parsed = ParsedDocument(
            format="txt",
            surfaces=[
                Surface(
                    id="heldout-surface",
                    source_type="BODY_TEXT",
                    location="原创合成句子",
                    text=case["text"],
                )
            ],
        )
        sentence_rows.append(inspect_case(case, parsed))
    for index, case in enumerate(documents):
        try:
            path = create_document(case)
            parsed = parse_document(path)
            row = inspect_case(case, parsed, path)
            row["document_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        except Exception as exc:
            row = dict(
                id=case["id"],
                format=case["format"],
                kind=case["kind"],
                target=case["target"],
                expected_risk=case["expected_risk"],
                candidate=False,
                confirmed_fail=False,
                actual_pass=False,
                rule_pass=False,
                coverage="UNAVAILABLE",
                diagnostics=[type(exc).__name__],
                evidence=[],
                statuses=[],
                ocr_regions=0,
            )
        document_rows.append(row)
        print(f"Hold-out document {index + 1}/{len(documents)}", flush=True)
    result = dict(
        version=1,
        dataset_sha256=digest,
        annotation_status=data["annotation_status"],
        protocol=data["protocol"],
        mode="SMOKE" if args.smoke else "FULL",
        sentence_metrics=summarize(sentence_rows),
        document_metrics=summarize(document_rows),
        sentence_cases=sentence_rows,
        document_cases=document_rows,
        interpretation="Separate frozen synthetic hold-out, not an external human blind study. REVIEW counts only as evidence candidate. PARTIAL/MANUAL/UNAVAILABLE never count as PASS. False PASS also reports per-rule PASS on positive cases.",
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(
        json.dumps(
            {k: result[k] for k in ("sentence_metrics", "document_metrics")},
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
