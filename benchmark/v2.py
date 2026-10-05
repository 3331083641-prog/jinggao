"""Reproducible synthetic candidate/confirmed verdict benchmarks. Preserve v1."""

import ast
import io
import json
import sys
import zipfile
from pathlib import Path
from time import perf_counter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
from app.schemas.domain import Rule, Surface, ParsedDocument
from app.detectors.engine import inspect_rule
from app.parsers.document import parse_document
from app.services.ocr import provider
from app.services.coverage import rule_coverage
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as SlideInches
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

DEST = ROOT / "tests/generated/benchmark-v2"


def metrics(rows, field):
    tp = sum(r["expected_risk"] and r[field] for r in rows)
    fp = sum(not r["expected_risk"] and r[field] for r in rows)
    fn = sum(r["expected_risk"] and not r[field] for r in rows)
    tn = sum(not r["expected_risk"] and not r[field] for r in rows)
    precision, recall = tp / max(tp + fp, 1), tp / max(tp + fn, 1)
    positives = sum(r["expected_risk"] for r in rows)
    false_pass = sum(r["expected_risk"] and r["actual_pass"] for r in rows)
    return dict(
        tp=tp,
        fp=fp,
        fn=fn,
        tn=tn,
        precision=precision,
        recall=recall,
        f1=2 * precision * recall / max(precision + recall, 1e-12),
        false_positive=fp,
        false_negative=fn,
        false_pass=false_pass,
        false_pass_rate=false_pass / max(positives, 1),
    )


def sentences():
    # Read v1 data through AST without running its evaluator or replacing results.json.
    tree = ast.parse((ROOT / "benchmark/evaluate.py").read_text(encoding="utf-8"))
    cases = next(
        ast.literal_eval(n.value)
        for n in tree.body
        if isinstance(n, ast.Assign)
        and any(isinstance(t, ast.Name) and t.id == "CASES" for t in n.targets)
    )
    cases += [
        ("school_name", "作者就读于清华。", True),
        ("school_name", "本研究由北大团队开展。", True),
        ("organization", "隶属中科院实验中心。", True),
        ("organization", "Affiliation: School of Synthetic Science.", True),
        (
            "organization",
            "Author affiliation: Department of Synthetic Engineering.",
            True,
        ),
        ("organization", "作者单位：合成重点实验室。", True),
        ("organization", "隶属合成科研研究中心。", True),
        ("telephone", "电话：0371-87654321", True),
        ("telephone", "联系电话：010-87654321", True),
        ("project_number", "Supported by NSFC 12345678.", True),
        ("author_name", "王样本负责本文撰写。", True),
        ("organization", "References: Example University Press, 2020.", False),
        ("school_name", "参考文献：清华团队的公开成果。", False),
        ("school_name", "引用资料来自北大，非作者归属。", False),
        ("school_name", "中科院研究所", False),
        ("school_name", "Department of Synthetic Science", False),
        ("contact", "参数编号为 1234567890。", False),
        ("person", "王方程是这里的算法名。", False),
        ("funding", "实验参数 NSFC 123 并非项目号。", False),
        ("organization", "General anonymous material without affiliations.", False),
    ]
    rows = []
    for target, text, expected in cases:
        rule = Rule(
            id=target,
            target=target,
            category="合成基准",
            description="实体候选",
            scope=["BODY_TEXT"],
        )
        parsed = ParsedDocument(
            format="txt",
            surfaces=[
                Surface(
                    id="sentence", source_type="BODY_TEXT", location="合成句", text=text
                )
            ],
        )
        findings = inspect_rule(rule, parsed, ["body"])
        evidence = [f for f in findings if f["surface_id"] == "sentence"]
        rows.append(
            dict(
                target=target,
                text=text,
                expected_risk=expected,
                candidate=bool(evidence),
                confirmed_fail=any(f["status"] == "FAIL" for f in evidence),
                actual_pass=any(f["status"] == "PASS" for f in findings),
                statuses=[f["status"] for f in findings],
            )
        )
    return rows


def font(size):
    # Only use installed font files; never copy/distribute Windows fonts.
    for p in [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]:
        if p.exists():
            return ImageFont.truetype(str(p), size)
    return ImageFont.load_default(size=size)


def picture(kind):
    image = Image.new("RGB", (1000, 350), "white")
    draw = ImageDraw.Draw(image)
    text = {
        "image_identity": "Author affiliation: Example University",
        "ai_watermark": "AI Generated",
        "normal_small_font": "Normal academic diagram: x = 1.35, F1 score 0.98",
        "garbled": "Replacement: \ufffd\ufffd\ufffd\ufffd",
        "missing_glyph": "Missing glyphs: □□□□",
    }.get(kind)
    if text:
        draw.text(
            (25, 120),
            text,
            fill="black",
            font=font(14 if kind == "normal_small_font" else 36),
        )
    else:
        draw.line((50, 290, 50, 40, 920, 40), fill="#444444", width=2)
        draw.line((80, 260, 250, 220, 450, 130, 820, 70), fill="#557788", width=3)
        draw.text((400, 300), "Synthetic chart", fill="black", font=font(18))
    return image


def fixture(fmt, kind):
    path = DEST / (kind + "." + fmt)
    is_image = kind in (
        "image_identity",
        "ai_watermark",
        "normal_small_font",
        "normal_chart",
        "garbled",
        "missing_glyph",
    )
    image = picture(kind) if is_image else None
    text = (
        "Author affiliation: Example University"
        if kind == "identity"
        else "Synthetic public document. x = 2 + 3."
    )
    if fmt == "pdf":
        c = canvas.Canvas(str(path), invariant=1)
        c.setAuthor("Synthetic Benchmark Author" if kind == "metadata" else "")
        if image:
            c.drawImage(ImageReader(image), 40, 350, width=500, height=175)
        else:
            c.drawString(40, 740, text)
        c.save()
    elif fmt == "docx":
        d = Document()
        d.core_properties.author = (
            "Synthetic Benchmark Author" if kind == "metadata" else ""
        )
        d.add_paragraph(text)
        if image:
            buffer = io.BytesIO()
            image.save(buffer, "PNG")
            buffer.seek(0)
            d.add_picture(buffer, width=Inches(5))
        if kind == "hidden_text":
            r = d.add_paragraph().add_run("Synthetic hidden author")
            r.font.hidden = True
        d.save(path)
        if kind in ("comment", "revision"):
            from lxml import etree

            with zipfile.ZipFile(path) as archive:
                parts = {p: archive.read(p) for p in archive.namelist()}
            ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            if kind == "comment":
                root = etree.Element("{" + ns + "}comments", nsmap={"w": ns})
                cm = etree.SubElement(
                    root,
                    "{" + ns + "}comment",
                    {"{" + ns + "}id": "0", "{" + ns + "}author": "Synthetic Reviewer"},
                )
                p = etree.SubElement(cm, "{" + ns + "}p")
                r = etree.SubElement(p, "{" + ns + "}r")
                etree.SubElement(r, "{" + ns + "}t").text = "Synthetic comment"
                parts["word/comments.xml"] = etree.tostring(root)
            else:
                root = etree.fromstring(parts["word/document.xml"])
                p = root.find(".//{" + ns + "}p")
                ins = etree.SubElement(
                    p,
                    "{" + ns + "}ins",
                    {"{" + ns + "}author": "Synthetic Editor", "{" + ns + "}id": "1"},
                )
                r = etree.SubElement(ins, "{" + ns + "}r")
                etree.SubElement(r, "{" + ns + "}t").text = "Synthetic revision content"
                parts["word/document.xml"] = etree.tostring(root)
            with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
                for name, data in parts.items():
                    archive.writestr(name, data)
    else:
        p = Presentation()
        p.core_properties.author = (
            "Synthetic Benchmark Author" if kind == "metadata" else ""
        )
        slide = p.slides.add_slide(p.slide_layouts[6])
        slide.shapes.add_textbox(
            SlideInches(0.5), SlideInches(0.5), SlideInches(8), SlideInches(1)
        ).text = text
        if image:
            buffer = io.BytesIO()
            image.save(buffer, "PNG")
            buffer.seek(0)
            slide.shapes.add_picture(
                buffer, SlideInches(0.5), SlideInches(2), width=SlideInches(8)
            )
        p.save(path)
    return path


def documents():
    DEST.mkdir(parents=True, exist_ok=True)
    rows = []
    for fmt in ("pdf", "docx", "pptx"):
        kinds = [
            "normal",
            "identity",
            "metadata",
            "image_identity",
            "ai_watermark",
            "normal_small_font",
            "normal_chart",
            "garbled",
            "missing_glyph",
        ]
        if fmt == "docx":
            kinds += ["hidden_text", "comment", "revision"]
        for kind in kinds:
            path = fixture(fmt, kind)
            parsed = parse_document(path)
            if parsed.image_jobs:
                provider.inspect(path, parsed, lambda *_: None)
            target, method, scope = ("school_name", "recognizer", ["BODY_TEXT"])
            expected = kind not in ("normal", "normal_small_font", "normal_chart")
            if kind == "metadata":
                target, method, scope = "metadata", "presence", ["METADATA"]
            elif kind in ("hidden_text", "comment", "revision"):
                target, method, scope = (
                    "hidden",
                    "presence",
                    [
                        {
                            "hidden_text": "HIDDEN_TEXT",
                            "comment": "COMMENT",
                            "revision": "REVISION",
                        }[kind]
                    ],
                )
            elif kind == "image_identity":
                scope = ["IMAGE_OCR"]
            elif kind == "ai_watermark":
                target, method, scope = "ai_marker", "ai_marker", ["IMAGE_OCR"]
            elif kind in (
                "normal_small_font",
                "normal_chart",
                "garbled",
                "missing_glyph",
            ):
                target, method, scope = "text_quality", "text_quality", ["IMAGE_OCR"]
            rule = Rule(
                id="benchmark",
                target=target,
                detection_method=method,
                scope=scope,
                category="合成文档",
                description=kind,
            )
            findings = inspect_rule(
                rule, parsed, ["body", "metadata", "hidden", "images"]
            )
            evidence = [
                f for f in findings if f["surface_id"] and f["status"] != "PASS"
            ]
            rows.append(
                dict(
                    format=fmt,
                    case=kind,
                    expected_risk=expected,
                    candidate=bool(evidence),
                    confirmed_fail=any(f["status"] == "FAIL" for f in evidence),
                    actual_pass=any(f["status"] == "PASS" for f in findings),
                    statuses=[f["status"] for f in findings],
                    ocr_regions=len(
                        [s for s in parsed.surfaces if s.source_type == "IMAGE_OCR"]
                    ),
                    coverage=rule_coverage(
                        rule, parsed, ["body", "metadata", "hidden", "images"], findings
                    )["status"],
                    diagnostics=parsed.warnings,
                )
            )
    return rows


def main():
    start = perf_counter()
    a, b = sentences(), documents()
    result = dict(
        version=2,
        dataset="56 synthetic sentences + 30 generated documents; no private uploads",
        boundary="Candidate metrics count evidence-backed FAIL+REVIEW; confirmed_fail metrics count only FAIL against seeded risk labels. Neither is real-world accuracy. Unknown is not PASS. OCR can normalize damaged glyphs: measured misses remain in results.",
        sentence=dict(
            count=len(a),
            candidate_metrics=metrics(a, "candidate"),
            confirmed_fail_metrics=metrics(a, "confirmed_fail"),
            cases=a,
        ),
        document=dict(
            count=len(b),
            candidate_metrics=metrics(b, "candidate"),
            confirmed_fail_metrics=metrics(b, "confirmed_fail"),
            cases=b,
        ),
        seconds=round(perf_counter() - start, 3),
    )
    (ROOT / "benchmark/results_v2.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        json.dumps(
            {
                k: {
                    "count": result[k]["count"],
                    "candidate": result[k]["candidate_metrics"],
                    "confirmed_fail": result[k]["confirmed_fail_metrics"],
                }
                for k in ("sentence", "document")
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
