import json
import time
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw
from pptx import Presentation
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from app.core import storage as db
from app.main import app
from app.schemas.domain import Rule, ParsedDocument, Surface
from app.services.cleanup import preview, create_copy, sha
from app.services.coverage import rule_coverage
from app.services.local_models import loopback_url
from app.services.rule_compiler import compile_rules
from app.services.run_diff import compare_runs
from app.services.vision import badge_regions
from app.detectors.engine import inspect_rule


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DATA", tmp_path)
    monkeypatch.setattr(db, "DB", tmp_path / "test.sqlite3")
    (tmp_path / "uploads").mkdir()
    (tmp_path / "reports").mkdir()
    with TestClient(app) as client:
        yield client


def wait(client, run):
    for _ in range(1500):
        value = client.get("/api/runs/" + run["id"]).json()
        if value["state"] not in ("QUEUED", "RUNNING"):
            assert value["state"] == "COMPLETED", value
            return value
        time.sleep(0.01)
    pytest.fail("Run timeout")


@pytest.mark.parametrize(
    "url",
    [
        "https://example.com/v1",
        "http://127.0.0.1@evil.invalid/v1",
        "http://localhost/?redirect=evil",
        "http://192.168.1.3/v1",
        "file:///model",
    ],
)
def test_provider_never_accepts_public_or_credential_url(url):
    with pytest.raises(ValueError):
        loopback_url(url)


def test_compiler_no_model_fallback(monkeypatch):
    monkeypatch.delenv("JINGGAO_LOCAL_LLM_BASE_URL", raising=False)
    monkeypatch.delenv("JINGGAO_LOCAL_LLM_MODEL", raising=False)
    result = compile_rules("正文不得出现学校名称。")
    assert result["compiler"]["mode"] == "DETERMINISTIC_FALLBACK"
    assert [r["target"] for r in result["rules"]] == ["school_name"]


class ProposedModel:
    configured = True

    def __init__(self, mode):
        self.mode = mode

    def complete(self, messages):
        rules = json.loads(messages[1]["content"])
        if self.mode == "invent":
            rules.append({"id": "email-invented", "original_text": "禁止邮箱"})
        elif self.mode == "broaden":
            rules[0]["scope"].append("METADATA")
        elif self.mode == "condition":
            rules[0]["condition"] = "原文没有的条件"
        else:
            for rule in rules:
                rule["normalized_requirement"] = "建议：依据原文核对学校名称"
                rule["confidence"] = 0.8
        return {"rules": rules}


@pytest.mark.parametrize("mode", ["invent", "broaden", "condition"])
def test_compiler_rejects_model_new_rules_or_authority(mode):
    result = compile_rules("正文不得出现学校名称。", ProposedModel(mode))
    assert result["compiler"]["mode"] == "DETERMINISTIC_FALLBACK"
    assert len(result["rules"]) == 1
    assert result["rules"][0]["scope"] == ["BODY_TEXT"]
    assert result["compiler_warnings"]


def test_grounded_local_semantic_proposal_is_advisory():
    result = compile_rules("正文不得出现学校名称。", ProposedModel("valid"))
    assert result["compiler"]["mode"] == "LOCAL_MODEL_VALIDATED"
    assert result["rules"][0]["needs_confirmation"]
    assert result["rules"][0]["target"] == "school_name"


@pytest.mark.parametrize(
    "clause",
    [
        "若采用匿名评审，正文不得出现学校名称。",
        "匿名评审材料正文及附录中不得出现学校名称，但参考文献中的机构名称不受此限制。",
        "除非允许实名登记，正文不得出现学校名称。",
    ],
)
def test_conditions_and_reference_exception_stay_grounded_manual(clause):
    rules = compile_rules(clause)["rules"]
    assert all(r["original_text"] in clause for r in rules)
    assert any(r["condition"] or r["exception"] for r in rules)
    parsed = ParsedDocument(
        format="txt",
        surfaces=[
            Surface(
                id="s", source_type="BODY_TEXT", location="正文", text="学校：合成大学"
            )
        ],
    )
    for rule in rules:
        findings = inspect_rule(Rule(**rule), parsed, ["body"])
        assert not any(f["status"] == "FAIL" for f in findings)


@pytest.mark.parametrize(
    "target,text",
    [
        ("school_name", "作者来自北大。"),
        ("school_name", "作者单位：清华。"),
        ("organization_name", "作者单位：中科院合成研究中心。"),
        ("school_name", "Author affiliation: Example University."),
        ("organization_name", "Department of Synthetic Studies"),
        ("organization_name", "本研究来自青岚重点实验室。"),
        ("telephone", "联系 010-87654321"),
        ("telephone", "电话 0371-12345678"),
        ("project_number", "NSFC 12345678"),
        ("author_name", "王合成负责本文撰写。"),
    ],
)
def test_layered_entities(target, text):
    rule = Rule(
        id="one",
        category="身份",
        target=target,
        description="唯一实体规则",
        scope=["BODY_TEXT"],
        source_rule_set_id="custom",
    )
    parsed = ParsedDocument(
        format="txt",
        surfaces=[Surface(id="s", source_type="BODY_TEXT", location="正文", text=text)],
    )
    findings = inspect_rule(rule, parsed, ["body"])
    assert any(f["surface_id"] == "s" for f in findings)
    if target == "author_name":
        assert all(f["status"] == "REVIEW" for f in findings)


def test_reference_affiliation_is_not_fail():
    rule = Rule(
        id="one",
        category="身份",
        target="school_name",
        description="学校",
        scope=["BODY_TEXT"],
        source_rule_set_id="custom",
    )
    parsed = ParsedDocument(
        format="txt",
        surfaces=[
            Surface(
                id="s",
                source_type="BODY_TEXT",
                location="参考文献",
                text="参考文献引用了 Example University 的成果。",
            )
        ],
    )
    assert all(f["status"] != "FAIL" for f in inspect_rule(rule, parsed, ["body"]))


def office_fixture(path):
    if path.suffix == ".docx":
        doc = Document()
        doc.add_paragraph("正文公式与引用保留 Example University。")
        doc.core_properties.author = "Synthetic Author"
        hidden = doc.add_paragraph().add_run("HIDDEN SYNTHETIC")
        hidden.font.hidden = True
        doc.save(path)
        with zipfile.ZipFile(path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        parts["word/document.xml"] = parts["word/document.xml"].replace(
            b"</w:body>",
            b'<w:p><w:ins w:id="1" w:author="Synthetic Reviewer" w:date="2026-01-01T00:00:00Z"><w:r><w:t>INSERT PRESERVED</w:t></w:r></w:ins><w:del w:id="2" w:author="Synthetic Reviewer"><w:r><w:delText>DELETE PRESERVED</w:delText></w:r></w:del></w:p></w:body>',
        )
        with zipfile.ZipFile(path, "w") as archive:
            for name, data in parts.items():
                archive.writestr(name, data)
    else:
        doc = Presentation()
        doc.slides.add_slide(doc.slide_layouts[6]).shapes.add_textbox(
            0, 0, 4000000, 500000
        ).text = "Visible Synthetic Content"
        doc.core_properties.author = "Synthetic Author"
        doc.save(path)


@pytest.mark.parametrize("fmt", ["docx", "pptx", "pdf"])
def test_cleanup_never_overwrites_and_reopens(tmp_path, fmt):
    path = tmp_path / ("synthetic." + fmt)
    if fmt == "pdf":
        c = canvas.Canvas(str(path))
        c.setAuthor("Synthetic Author")
        c.drawString(80, 700, "Original body intact")
        c.save()
    else:
        office_fixture(path)
    before = sha(path)
    operations = ["metadata"] + (
        ["hidden_text", "revision_metadata"] if fmt == "docx" else []
    )
    plan = preview(path, operations)
    assert plan["change_count"] > 0
    target, name, _ = create_copy(path, path.name, tmp_path, operations, before)
    assert sha(path) == before and target != path and ".cleaned." in name
    if fmt == "pdf":
        assert PdfReader(target).metadata is None
        assert "Original body intact" in PdfReader(target).pages[0].extract_text()
    else:
        with zipfile.ZipFile(target) as archive:
            assert not archive.testzip()
            if fmt == "docx":
                content = archive.read("word/document.xml")
                assert b"HIDDEN SYNTHETIC" not in content
                assert b"INSERT PRESERVED" in content and b"DELETE PRESERVED" in content
                assert b"Synthetic Reviewer" not in content


def test_cleanup_api_same_snapshot_run_n_plus_one(client, tmp_path):
    source = tmp_path / "synthetic.docx"
    office_fixture(source)
    before = source.read_bytes()
    payload = {
        "name": "只检查属性",
        "rules": [
            {
                "id": "metadata",
                "category": "属性",
                "target": "metadata",
                "description": "清除作者元数据",
                "detection_method": "presence",
                "scope": ["METADATA"],
            }
        ],
    }
    rules = client.post("/api/rulesets", json=payload).json()
    original = wait(
        client,
        client.post(
            "/api/generate",
            data={"ruleset_id": rules["id"]},
            files={"file": (source.name, before)},
        ).json(),
    )
    assert original["coverage_matrix"][0]["status"] == "VERIFIED"
    plan = client.post(
        "/api/runs/" + original["id"] + "/cleanup-preview",
        json={"operations": ["metadata"]},
    ).json()
    next_run = wait(
        client,
        client.post(
            "/api/runs/" + original["id"] + "/cleanup-copy",
            json={"preview_token": plan["preview_token"]},
        ).json(),
    )
    assert next_run["version"] == original["version"] + 1
    assert next_run["ruleset_snapshot"] == original["ruleset_snapshot"]
    assert source.read_bytes() == before
    assert next_run["document"]["name"].endswith(".cleaned.docx")
    assert next_run["counts"]["FAIL"] == 0
    diff = client.get("/api/runs/" + next_run["id"] + "/diff").json()
    assert diff["comparable"] and any(
        pair["category"] == "FAIL_TO_PASS" for pair in diff["pairs"]
    )
    assert any(
        pair["before"]["evidence"] == "Synthetic Author" for pair in diff["pairs"]
    )
    report = client.get("/api/runs/" + next_run["id"] + "/report.pdf")
    assert report.content.startswith(b"%PDF")
    assert (
        client.post(
            "/api/runs/" + original["id"] + "/cleanup-copy",
            json={"preview_token": plan["preview_token"]},
        ).status_code
        == 409
    )


def test_stale_preview_rejects_copy(tmp_path):
    source = tmp_path / "file.docx"
    office_fixture(source)
    before = sha(source)
    doc = Document(source)
    doc.add_paragraph("later change")
    doc.save(source)
    with pytest.raises(ValueError):
        create_copy(source, source.name, tmp_path, ["metadata"], before)


def test_coverage_is_unavailable_before_execution():
    r = Rule(
        id="a",
        category="x",
        target="metadata",
        description="属性",
        detection_method="presence",
        scope=["METADATA"],
    )
    p = ParsedDocument(
        format="pdf",
        surfaces=[
            Surface(id="m", source_type="METADATA", location="属性", text="Author")
        ],
    )
    assert (
        rule_coverage(r, p, ["metadata"], [], executed=False)["status"] == "UNAVAILABLE"
    )
    assert rule_coverage(r, p, [], [], executed=True)["status"] == "PARTIAL"
    r.detection_method = "manual"
    assert rule_coverage(r, p, ["metadata"], [])["status"] == "MANUAL"


def test_geometric_candidate_is_not_every_image():
    empty = Image.new("RGB", (600, 400), "white")
    assert badge_regions(empty) == []
    badge = empty.copy()
    draw = ImageDraw.Draw(badge)
    draw.ellipse((80, 80, 160, 160), outline="#285738", width=5)
    draw.ellipse((95, 95, 145, 145), outline="#285738", width=2)
    assert badge_regions(badge)


def test_visual_provider_only_runs_for_enabled_rule(client, tmp_path, monkeypatch):
    calls = []
    monkeypatch.setattr(
        "app.services.vision.prepare_visual", lambda *args: calls.append(args)
    )
    source = tmp_path / "normal.docx"
    office_fixture(source)
    rules = client.post(
        "/api/rulesets",
        json={
            "name": "非视觉",
            "rules": [
                {
                    "id": "literal",
                    "category": "文字",
                    "target": "literal",
                    "description": "禁止词",
                    "detection_method": "literal",
                    "scope": ["BODY_TEXT"],
                    "parameters": {"text": "SYNTHETIC_BAD"},
                }
            ],
        },
    ).json()
    wait(
        client,
        client.post(
            "/api/generate",
            data={"ruleset_id": rules["id"]},
            files={"file": (source.name, source.read_bytes())},
        ).json(),
    )
    assert calls == []


def test_diff_disappeared_is_not_false_pass():
    before = {
        "id": "old",
        "ruleset_snapshot": {"rules": []},
        "scopes": ["body"],
        "state": "COMPLETED",
        "counts": {"FAIL": 1, "REVIEW": 0, "PASS": 0},
        "findings": [
            {
                "rule_id": "a",
                "source_type": "BODY_TEXT",
                "page": None,
                "evidence": "x",
                "status": "FAIL",
            }
        ],
    }
    after = {
        **before,
        "id": "new",
        "counts": {"FAIL": 0, "REVIEW": 0, "PASS": 0},
        "findings": [],
        "coverage_matrix": [{"rule_id": "a", "status": "PARTIAL"}],
    }
    assert compare_runs(before, after)["pairs"][0]["category"] == "disappeared"


@pytest.mark.parametrize(
    "clause", ["可以使用公开数据。", "仅限正文使用匿名表述。", "正文不可出现学校名称。"]
)
def test_permission_and_limited_requirements_preserved(clause):
    result = compile_rules(clause)
    assert all(r["original_text"] in clause for r in result["rules"])
    assert all(r["needs_confirmation"] for r in result["rules"])


def test_reference_explicit_not_author_is_not_fail():
    r = Rule(
        id="s",
        target="school_name",
        category="学校",
        description="仅学校",
        scope=["BODY_TEXT"],
        source_rule_set_id="custom",
    )
    p = ParsedDocument(
        format="txt",
        surfaces=[
            Surface(
                id="s",
                source_type="BODY_TEXT",
                location="引用",
                text="引用资料来自北大，非作者归属。",
            )
        ],
    )
    assert not any(f["status"] == "FAIL" for f in inspect_rule(r, p, ["body"]))


def test_model_cannot_drop_reference_exception():
    from app.services.rule_compiler.schema_validator import validate_proposal

    text = "正文不得出现学校名称，但参考文献除外。"
    baseline = compile_rules(text)["rules"][0]
    with pytest.raises(ValueError):
        validate_proposal(
            {"original_text": baseline["original_text"], "exception": ""},
            baseline,
            text,
        )


@pytest.mark.parametrize(
    "target,text",
    [("ai_marker", "Al Generated"), ("text_quality", "Normal diagram x = 1.2")],
)
def test_ocr_visual_limits_do_not_false_pass(target, text):
    r = Rule(
        id="r",
        target=target,
        detection_method=target,
        category="图片",
        description=target,
        scope=["IMAGE_OCR"],
    )
    p = ParsedDocument(
        format="docx",
        image_jobs=[{"part": "image.png"}],
        surfaces=[
            Surface(
                id="ocr",
                source_type="IMAGE_OCR",
                location="图片",
                text=text,
                confidence=0.99,
            )
        ],
    )
    findings = inspect_rule(r, p, ["images"])
    assert not any(f["status"] == "PASS" for f in findings)
    assert rule_coverage(r, p, ["images"], findings)["status"] == "PARTIAL"
    if target == "ai_marker":
        assert findings[0]["status"] == "REVIEW" and findings[0]["evidence"] == text
    else:
        assert findings == []  # No confidence-only review.


def test_cleanup_comments_custom_props_and_links(tmp_path):
    from lxml import etree

    p = tmp_path / "synthetic.docx"
    office_fixture(p)
    with zipfile.ZipFile(p) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["word/comments.xml"] = (
        b'<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:comment w:id="0"><w:p><w:r><w:t>Private synthetic comment</w:t></w:r></w:p></w:comment></w:comments>'
    )
    parts["docProps/custom.xml"] = (
        b'<Properties><property name="Synthetic"><value>Private synthetic</value></property></Properties>'
    )
    root = etree.fromstring(parts["word/_rels/document.xml.rels"])
    ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    etree.SubElement(
        root,
        "{" + ns + "}Relationship",
        Id="rIdSynthetic",
        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        Target="file:///synthetic-private-location",
        TargetMode="External",
    )
    parts["word/_rels/document.xml.rels"] = etree.tostring(root)
    with zipfile.ZipFile(p, "w") as z:
        for name, data in parts.items():
            z.writestr(name, data)
    before = sha(p)
    target, _, _ = create_copy(
        p,
        p.name,
        tmp_path,
        ["comments", "custom_properties", "dangerous_links"],
        before,
    )
    assert sha(p) == before
    with zipfile.ZipFile(target) as z:
        assert "word/comments.xml" not in z.namelist()
        assert len(etree.fromstring(z.read("docProps/custom.xml"))) == 0
        assert b"file:///" not in z.read("word/_rels/document.xml.rels")
        assert b"INSERT PRESERVED" in z.read("word/document.xml")


def test_cleanup_signed_or_form_pdf_rejected(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject

    p = tmp_path / "form.pdf"
    w = PdfWriter()
    w.add_blank_page(width=100, height=100)
    w._root_object[NameObject("/AcroForm")] = DictionaryObject()
    with p.open("wb") as out:
        w.write(out)
    with pytest.raises(ValueError):
        preview(p, ["metadata"])


def test_no_unverified_evidence_crop(tmp_path):
    from app.reports.evidence_image import evidence_image

    f = {"bbox": [0, 0, 100, 100], "page": 1, "surface_id": "missing"}
    assert (
        evidence_image(
            {
                "format": "pdf",
                "path": str(tmp_path / "missing.pdf"),
                "parsed": {"surfaces": []},
            },
            f,
        )
        is None
    )


def test_visual_rule_gets_candidate_not_fail(tmp_path):
    from app.parsers.document import parse_document
    from app.services.vision import prepare_visual
    import io

    p = tmp_path / "badge.docx"
    d = Document()
    d.add_paragraph("Synthetic diagram")
    im = Image.new("RGB", (600, 400), "white")
    dr = ImageDraw.Draw(im)
    dr.ellipse((80, 80, 160, 160), outline="#285738", width=5)
    dr.ellipse((95, 95, 145, 145), outline="#285738", width=2)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    d.add_picture(buf)
    d.save(p)
    parsed = parse_document(p)
    prepare_visual(p, parsed)
    assert parsed.visual_candidates
    r = Rule(
        id="logo",
        target="logo",
        category="图形",
        description="仅Logo",
        scope=["IMAGE_OCR", "LOGO"],
        detection_method="visual",
    )
    findings = inspect_rule(r, parsed, ["images"])
    assert findings and all(f["status"] == "REVIEW" for f in findings)
    assert rule_coverage(r, parsed, ["images"], findings)["status"] == "PARTIAL"


def test_diff_different_snapshot_never_claims_transition_pass():
    f = {
        "rule_id": "a",
        "source_type": "BODY_TEXT",
        "page": 1,
        "evidence": "old",
        "status": "FAIL",
    }
    before = {
        "id": "old",
        "ruleset_snapshot": {"rules": [{"id": "a"}]},
        "scopes": ["body"],
        "state": "COMPLETED",
        "counts": {"FAIL": 1, "REVIEW": 0, "PASS": 0},
        "findings": [f],
    }
    after = {
        **before,
        "id": "new",
        "ruleset_snapshot": {"rules": [{"id": "changed"}]},
        "findings": [{**f, "status": "PASS"}],
        "coverage_matrix": [{"rule_id": "a", "status": "VERIFIED"}],
    }
    result = compare_runs(before, after)
    assert not result["comparable"] and result["pairs"][0]["category"] == "disappeared"


def test_bad_local_model_address_falls_back_without_network(monkeypatch):
    monkeypatch.setenv("JINGGAO_LOCAL_LLM_BASE_URL", "https://example.invalid/v1")
    monkeypatch.setenv("JINGGAO_LOCAL_LLM_MODEL", "synthetic")
    result = compile_rules("正文不得出现学校名称。")
    assert result["compiler"]["mode"] == "DETERMINISTIC_FALLBACK"
    assert result["compiler_warnings"]


def test_real_pdf_crop_from_verified_coordinates(tmp_path):
    from app.parsers.document import parse_document
    from app.reports.evidence_image import evidence_image

    p = tmp_path / "evidence.pdf"
    c = canvas.Canvas(str(p))
    c.drawString(60, 700, "Synthetic bounded evidence")
    c.save()
    parsed = parse_document(p)
    s = next(s for s in parsed.surfaces if s.source_type == "BODY_TEXT")
    doc = {"format": "pdf", "path": str(p), "parsed": parsed.model_dump()}
    f = {"page": s.page, "bbox": s.bbox, "surface_id": s.id}
    assert evidence_image(doc, f) is not None
